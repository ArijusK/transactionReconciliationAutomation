import pandas as pd
import pytest
import requests

from openpyxl import load_workbook
from unittest.mock import patch

import src.main
import src.report


@pytest.fixture
def report_path(tmp_path, monkeypatch):
    monkeypatch.setattr(src.report, "OUTPUT_DIR", tmp_path)

    with patch(
        "src.exchange_rates.requests.get",
        side_effect=requests.ConnectionError()
    ):
        assert src.main.run() == 0

    return tmp_path / "reconciliationReport.xlsx"


def test_missing_sheets_only_show_the_side_with_the_transaction(report_path):
    missing_from_bank = pd.read_excel(
        report_path,
        sheet_name="Missing From Bank"
    )

    missing_from_internal = pd.read_excel(
        report_path,
        sheet_name="Missing From Internal"
    )

    assert list(missing_from_bank.columns) == [
        "transaction_id",
        "date_internal",
        "currency_internal",
        "amount_internal",
        "status_internal",
    ]

    assert list(missing_from_internal.columns) == [
        "transaction_id",
        "date_bank",
        "currency_bank",
        "amount_bank",
        "status_bank",
    ]


def test_matched_transactions_are_counted_and_listed(report_path):
    summary = pd.read_excel(
        report_path,
        sheet_name="Summary"
    ).set_index("Metric")["Count"]

    matched = pd.read_excel(
        report_path,
        sheet_name="Matched By ID"
    )

    assert summary["Matched by ID"] == 5
    assert len(matched) == 5
    assert "_merge" not in matched.columns


def test_amounts_are_shown_with_two_decimals(report_path):
    worksheet = load_workbook(report_path)["Amount Mismatches"]
    headers = [cell.value for cell in worksheet[1]]

    for column in [
        "amount_internal",
        "amount_bank",
        "bank_amount_converted",
        "difference",
    ]:
        cell = worksheet.cell(row=2, column=headers.index(column) + 1)

        assert cell.number_format == "#,##0.00"
